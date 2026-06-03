# -*- coding: utf-8 -*-
"""v3 icerigini BKM Kitap sablonuna (kendi markali layout/tema) aktarir."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

TPL=r"C:\Users\fikri.eren\Desktop\Sunum.pptx"
OUT=r"C:\Users\fikri.eren\Desktop\Sunum_Egitim.pptx"
ICOW=os.path.join(os.path.dirname(__file__),"_icons")
ICOR=os.path.join(os.path.dirname(__file__),"_icons_red")

# BKM marka paleti — kirmizi baskin + gri tonlar (markaya uyumlu)
RED=RGBColor(0xE3,0x06,0x22); DRED=RGBColor(0xA6,0x00,0x1A); ROSE=RGBColor(0xC0,0x14,0x2B)
CHAR=RGBColor(0x2B,0x2B,0x2B); GREY=RGBColor(0x6E,0x6E,0x6E); LGREY=RGBColor(0xF2,0xF2,0xF2)
MGREY=RGBColor(0x9A,0x9A,0x9A); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x33,0x33,0x33)
# stack/kart icin markaya uyumlu kirmizi-gri kademe seti
SET=[RED,ROSE,RGBColor(0x7A,0x10,0x20),RGBColor(0x4A,0x4A,0x4A),RGBColor(0x8C,0x8C,0x8C)]

import copy as _copy
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
pr=Presentation(TPL); W,Hh=pr.slide_width,pr.slide_height
srcpr=Presentation(TPL)  # orijinal slaytlar (arka plan kopyalamak icin)
TITLE_BG=srcpr.slides[0]      # ikon-desenli kirmizi bantli zemin
CONTENT_BG=srcpr.slides[1]    # baslik bandi + logo'lu zemin
def L(name):
    for l in pr.slide_masters[0].slide_layouts:
        if l.name==name: return l
    return pr.slide_masters[0].slide_layouts[6]
def copy_bg(src,dst):
    csrc=src._element.find(qn('p:cSld')); bg=csrc.find(qn('p:bg'))
    if bg is None: return
    bg2=_copy.deepcopy(bg)
    for blip in bg2.iter(qn('a:blip')):
        rid=blip.get(qn('r:embed'))
        if not rid: continue
        img=src.part.related_part(rid)
        nrid=dst.part.relate_to(img,RT.IMAGE)
        blip.set(qn('r:embed'),nrid)
    cdst=dst._element.find(qn('p:cSld')); cdst.insert(0,bg2)
# ornek slaytlari sil
sld=pr.slides._sldIdLst
for sid in list(sld):
    try: pr.part.drop_rel(sid.get(qn('r:id')))
    except: pass
    sld.remove(sid)
def add(name):
    sl=pr.slides.add_slide(L(name))
    copy_bg(TITLE_BG if name in ("Başlık Slaydı","Bölüm Üst Bilgisi") else CONTENT_BG, sl)
    return sl
def setph(sl,idx,text):
    for ph in sl.placeholders:
        if ph.placeholder_format.idx==idx: ph.text=text; return ph
    return None
def tb(sl,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sp=1.0):
    bx=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=bx.text_frame
    tf.word_wrap=True;tf.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,None,None,None)]
    first=True
    for txt,sz,bd,col in runs:
        p=tf.paragraphs[0] if first else tf.add_paragraph();first=False
        p.alignment=align;p.line_spacing=sp
        r=p.add_run();r.text=txt;r.font.size=Pt(sz or 13);r.font.bold=bool(bd)
        r.font.name="Calibri";r.font.color.rgb=col or INK
    return bx
def rrect(sl,x,y,w,h,fill,line=None,rad=True,lw=1):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=fill
    if line: sh.line.color.rgb=line;sh.line.width=Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit=False;return sh
def circ(sl,x,y,d,bg,icon):
    c=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    c.fill.solid();c.fill.fore_color.rgb=bg;c.line.fill.background();c.shadow.inherit=False
    ip=d*0.55;sl.shapes.add_picture(os.path.join(ICOW,icon+".png"),
        Inches(x+(d-ip)/2),Inches(y+(d-ip)/2),Inches(ip),Inches(ip))
def card(sl,x,y,w,h,top=None):
    rrect(sl,x,y,w,h,WHITE,RGBColor(0xE0,0xE0,0xE0))
    if top: rrect(sl,x,y,w,0.07,top,rad=False)
def sig(sl):  # logo solda; imza sag altta (band/logo'ya dokunmaz)
    tb(sl,9.3,6.95,3.4,0.3,[("© Fikri Eren 2026",9,False,MGREY)],align=PP_ALIGN.RIGHT)

# ---------- 1 KAPAK (markali Baslik Slaydi) ----------
s=add("Başlık Slaydı")
setph(s,0,"Claude'u Süper Güce Dönüştürmek")
setph(s,1,"Ham model · handikaplar · skill / agent / hook / kural katmanının kattığı süper güçler")

# ---------- 2 GUNDEM ----------
s=add("Yalnızca Başlık"); setph(s,0,"Yolculuk altı durakta")
ag=[("cpu","Claude nedir?","Ne yapar, ortamlar, neyi iyi yapar."),
("alert-triangle","Handikaplar","Ham haliyle 6 büyük zaaf."),
("zap","Süper güçler","Hafıza, hook, skill, agent, kural."),
("layers","Çözüm haritası","Hangi zaafı hangi yapı kapattı."),
("workflow","Önce / Sonra","Somut: günaydın deyince ne olur."),
("rocket","Sonuç","Akıl Claude'da değil, yapıda.")]
for i,(ic,h,d) in enumerate(ag):
    x=0.6+(i%3)*4.15;y=1.5+(i//3)*2.15;col=SET[i%len(SET)]
    card(s,x,y,3.9,1.95,col);circ(s,x+0.25,y+0.22,0.58,RED,ic)
    tb(s,x+0.25,y+0.9,3.4,0.4,[(h,14,True,CHAR)])
    tb(s,x+0.25,y+1.28,3.45,0.6,[(d,11,False,GREY)])
sig(s)

# ---------- 3 DIVIDER 01 ----------
s=add("Bölüm Üst Bilgisi"); setph(s,0,"Claude nedir, ne yapar?"); setph(s,1,"Bölüm 1 — temeli netleştirelim"); sig(s)

# ---------- 4 CLAUDE NEDIR (Karsilastirma) ----------
s=add("Karşılaştırma"); setph(s,0,"Claude'un iki yüzü")
setph(s,1,"Genel Claude (sohbet)"); setph(s,3,"Claude Code (asistan)")
setph(s,2,"Soru-cevap, yazım, analiz. Kod örneği üretir — ama projene doğrudan dokunamaz.")
setph(s,4,"Dosyaları okur/yazar/çalıştırır · git commit · test koşar · her adımda onay ister."); sig(s)

# ---------- 5 ORTAMLAR ----------
s=add("Yalnızca Başlık"); setph(s,0,"Nerede kullanılır?")
ort=[("message-circle","Chat (claude.ai)","Düşünme, taslak, doküman. git/hook yok."),
("users","Cowork","Ajan görevleri. Skill + MCP + plugin."),
("monitor","Desktop / Terminal","Tam güç: dosya + git + hook + agent.")]
for i,(ic,h,d) in enumerate(ort):
    x=0.6+i*4.15;card(s,x,1.6,3.9,1.95,SET[i]);circ(s,x+0.25,1.82,0.58,RED,ic)
    tb(s,x+0.25,2.5,3.4,0.4,[(h,13.5,True,CHAR)]);tb(s,x+0.25,2.9,3.45,0.6,[(d,11,False,GREY)])
rrect(s,0.6,3.9,12.1,0.85,LGREY,RED,lw=1.5)
tb(s,0.9,3.9,11.5,0.85,[("Ama — kutudan çıktığı HAM haliyle ciddi handikapları var.",13,True,DRED)],anchor=MSO_ANCHOR.MIDDLE)
sig(s)

# ---------- 6 DIVIDER 02 ----------
s=add("Bölüm Üst Bilgisi"); setph(s,0,"Ham Claude'un handikapları"); setph(s,1,"Bölüm 2 — güçlü ama disiplinsiz"); sig(s)

# ---------- 7 STAT donut (kirmizi) ----------
s=add("Yalnızca Başlık"); setph(s,0,"En büyük zaaf: hafızasızlık")
cd=CategoryChartData();cd.categories=["kayıp"];cd.add_series(" ",(100,))
gf=s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT,Inches(0.7),Inches(1.6),Inches(3.7),Inches(3.7),cd).chart
gf.has_legend=False;gf.has_title=False;gf.plots[0].has_data_labels=False
gf.series[0].points[0].format.fill.solid();gf.series[0].points[0].format.fill.fore_color.rgb=RED
tb(s,0.7,2.85,3.7,0.8,[("%100",38,True,RED)],align=PP_ALIGN.CENTER)
tb(s,0.7,3.55,3.7,0.4,[("hafıza kaybı",12,False,GREY)],align=PP_ALIGN.CENTER)
tb(s,4.9,1.9,7.6,2.0,[("Her yeni oturumda TAM hafıza kaybı.",19,True,DRED),
 ("Claude dünü, kararları, talimatları hatırlamaz — her sabah hafızası silinmiş bir stajyer.",15,False,INK)],sp=1.1)
tb(s,4.9,4.1,7.6,0.6,[("Sonuç: tekrar anlatım · tutarsızlık · 'bir gün unutursun → bug'.",12,False,GREY)]); sig(s)

# ---------- 8 6 HANDIKAP ----------
s=add("Yalnızca Başlık"); setph(s,0,"Altı handikap")
hd=[("1 · Hafızasızlık","Her oturum sıfırdan. Dünü hatırlamaz."),
("2 · Tutarsızlık","Aynı işi her sefer farklı yapar."),("3 · Disiplinsizlik","Commit/test/plan'ı uygulamaz."),
("4 · Domain körlüğü","Senin işini bilmez; uydurabilir."),("5 · Riskli işlem","Dosya siler, sır sızdırır."),
("6 · Sahte 'bitti'","Çalıştırmadan 'tamam' der.")]
for i,(h,d) in enumerate(hd):
    x=0.6+(i%3)*4.15;y=1.5+(i//3)*2.15;card(s,x,y,3.9,1.95);rrect(s,x,y,0.09,1.95,RED,rad=False)
    s.shapes.add_picture(os.path.join(ICOR,"alert-triangle.png"),Inches(x+0.25),Inches(y+0.25),Inches(0.32),Inches(0.32))
    tb(s,x+0.7,y+0.2,3,0.4,[(h,13,True,CHAR)]);tb(s,x+0.25,y+0.78,3.45,0.9,[(d,11,False,GREY)])
sig(s)

# ---------- 9 DIVIDER 03 ----------
s=add("Bölüm Üst Bilgisi"); setph(s,0,"Bizim katman: süper güçler"); setph(s,1,"Bölüm 3 — ham model + .claude/ yapısı"); sig(s)

# ---------- 10 MIMARI STACK (kirmizi-gri kademe) ----------
s=add("Yalnızca Başlık"); setph(s,0,"Beş bileşenli yapı")
st=[("brain","Hafıza katmanı","CLAUDE.md + rules + journal + TODO"),
("cpu","Hook'lar","olay tetikli — unutsa bile mekanik çalışır"),
("package","Skills","çok adımlı iş reçeteleri"),
("bot","Agent'lar","izole uzmanlar + model katmanı"),
("scroll-text","Kurallar","14 evrensel kalıcı davranış")]
for i,(ic,h,d) in enumerate(st):
    x=0.6+i*0.22;y=1.55+i*0.84;w=10.8-i*0.22
    rrect(s,x,y,w,0.72,SET[i])
    s.shapes.add_picture(os.path.join(ICOW,ic+".png"),Inches(x+0.22),Inches(y+0.19),Inches(0.34),Inches(0.34))
    tb(s,x+0.72,y,3,0.72,[(h,13,True,WHITE)],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,x+3.5,y,w-3.7,0.72,[(d,11,False,WHITE)],anchor=MSO_ANCHOR.MIDDLE)
sig(s)

# ---------- 11 HOOK FLOW ----------
s=add("Yalnızca Başlık"); setph(s,0,"Hafıza dışsallaştırılır")
fn=[("brain","Yeni oturum","boş başlar",LGREY,GREY),
("cpu","session-start hook","git+TODO+journal",RGBColor(0xFC,0xE9,0xEC),RED),
("circle-check","bağlam yüklü","'nerede kaldık'",RGBColor(0xFC,0xE9,0xEC),RED)]
for i,(ic,h,d,bg,ln) in enumerate(fn):
    x=0.7+i*4.3;rrect(s,x,1.55,3.6,1.25,bg,ln,lw=1.5)
    s.shapes.add_picture(os.path.join(ICOR,ic+".png"),Inches(x+1.55),Inches(1.7),Inches(0.48),Inches(0.48))
    tb(s,x,2.22,3.6,0.35,[(h,12,True,CHAR)],align=PP_ALIGN.CENTER)
    tb(s,x,2.54,3.6,0.3,[(d,10,False,GREY)],align=PP_ALIGN.CENTER)
    if i<2: tb(s,x+3.6,1.55,0.7,1.25,[("➜",24,True,RED)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
hk=[("terminal","session-start","Açılışta otomatik durum özeti."),
("shield-check","pre-commit","Şifre/riskli desen → bloklar."),
("book-open","post-commit","Her commit → günlüğe kayıt.")]
for i,(ic,h,d) in enumerate(hk):
    x=0.6+i*4.15;y=3.2;card(s,x,y,3.9,1.6,SET[i]);circ(s,x+0.25,y+0.22,0.52,RED,ic)
    tb(s,x+0.9,y+0.28,2.8,0.4,[(h,12.5,True,CHAR)]);tb(s,x+0.25,y+0.82,3.45,0.7,[(d,10.5,False,GREY)])
sig(s)

# ---------- 12 DIVIDER 04 ----------
s=add("Bölüm Üst Bilgisi"); setph(s,0,"Handikap → nasıl güçlendirdik"); setph(s,1,"Bölüm 4 — eşleme haritası"); sig(s)

# ---------- 13 HARITA TABLO ----------
s=add("Yalnızca Başlık"); setph(s,0,"Her zaaf bir yapı ile kapandı")
rows=[("Ham zaaf","Çözüm (yapı)","Sonuç"),
("1 Hafızasızlık","3 katman hafıza + session-start hook","Her oturum dünü hatırlar"),
("2 Tutarsızlık","Kurallar + Skills (reçete)","Her sefer aynı doğru sonuç"),
("3 Disiplinsizlik","commit/plan/test kuralları + hook","Disiplin mekanik, atlanamaz"),
("4 Domain körlüğü","Proje-özel skill + kural + ADR","Senin işini bilir, uydurmaz"),
("5 Riskli işlem","pre-commit + before-major-change","Sır sızmaz, riskli işlem onaylı"),
("6 Sahte 'bitti'","verify + test-discipline + denetçi agent","Çalıştırmadan 'tamam' yok")]
t=s.shapes.add_table(len(rows),3,Inches(0.6),Inches(1.55),Inches(12.1),Inches(4.2)).table
t.columns[0].width=Inches(3.0);t.columns[1].width=Inches(5.0);t.columns[2].width=Inches(4.1)
for r,row in enumerate(rows):
    for c,val in enumerate(row):
        cell=t.cell(r,c);cell.text=val
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.size=Pt(11.5);run.font.name="Calibri";run.font.bold=(r==0)
                run.font.color.rgb=WHITE if r==0 else INK
        cell.fill.solid();cell.fill.fore_color.rgb=RED if r==0 else (LGREY if r%2 else WHITE)
sig(s)

# ---------- 14 ONCE/SONRA + bar ----------
s=add("Yalnızca Başlık"); setph(s,0,"Tek kelime: günaydın")
card(s,0.6,1.55,5.9,1.6,GREY);tb(s,0.85,1.68,5.4,0.3,[("ÖNCE · HAM",11,True,GREY)])
tb(s,0.85,2.02,5.4,1.0,[("Claude: 'Merhaba! Nasıl yardımcı olabilirim?' → Dünü bilmez. Baştan anlatırsın.",12,False,INK)])
card(s,6.8,1.55,5.9,1.6,RED);tb(s,7.05,1.68,5.4,0.3,[("SONRA · YAPI KURULU",11,True,RED)])
tb(s,7.05,2.02,5.4,1.1,[("Claude: 'Dün giriş bitti. Bugün şifre sıfırlama (Faz 0). 3 uncommitted. Başlayalım mı?' → session-start + handoff okudu.",12,False,INK)])
cd2=CategoryChartData();cd2.categories=["Önce","Sonra"];cd2.add_series("dk",(10,0.3))
bc=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(0.6),Inches(3.4),Inches(5.6),Inches(2.4),cd2).chart
bc.has_legend=False;bc.has_title=False;bc.plots[0].has_data_labels=True
bc.plots[0].vary_by_categories=False
bc.series[0].format.fill.solid();bc.series[0].format.fill.fore_color.rgb=RED
rrect(s,6.8,3.8,5.9,1.6,LGREY,RED,lw=1.5)
tb(s,7.05,3.8,5.5,1.6,[("\"günaydın\" sihirli değil. ",13,True,DRED),
 ("Arkada git + TODO + günlük + özet — hepsi senin kurduğun yapı sayesinde.",13,False,INK)],anchor=MSO_ANCHOR.MIDDLE)
sig(s)

# ---------- 15 TETIK TABLO ----------
s=add("Yalnızca Başlık"); setph(s,0,"Söz → yapı → eylem")
tr=[("Sen dersin","Hangi yapı sayesinde","Ne yapar"),
("günaydın","session-start hook + handoff","git/TODO/journal okur, özetler"),
("iyi geceler","session-handoff skill","günlüğe yazar, TODO işaretler, commit'ler"),
("dosyaları böl","commit-splitter agent","15+ dosyayı bucket'lara ayırır"),
("hangi seçenek?","llm-council skill","5 danışman + sentez"),
("(commit'te şifre)","pre-commit hook","sızıntıyı yakalar, bloklar")]
t2=s.shapes.add_table(len(tr),3,Inches(0.6),Inches(1.55),Inches(12.1),Inches(3.5)).table
t2.columns[0].width=Inches(3.2);t2.columns[1].width=Inches(4.6);t2.columns[2].width=Inches(4.3)
for r,row in enumerate(tr):
    for c,val in enumerate(row):
        cell=t2.cell(r,c);cell.text=val
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.size=Pt(12);run.font.name="Calibri";run.font.bold=(r==0)
                run.font.color.rgb=WHITE if r==0 else INK
        cell.fill.solid();cell.fill.fore_color.rgb=RED if r==0 else (LGREY if r%2 else WHITE)
rrect(s,0.6,5.3,12.1,0.85,LGREY,RED,lw=1.5)
tb(s,0.9,5.3,11.5,0.85,[("Her kolaylık = bir tetik söz + senin kurduğun bir yapı. Yapıyı sil → kolaylık kaybolur.",13.5,True,DRED)],anchor=MSO_ANCHOR.MIDDLE)
sig(s)

# ---------- 16 KAPANIS (markali Baslik Slaydi) ----------
s=add("Başlık Slaydı")
setph(s,0,"Akıl Claude'da değil — kurduğun yapıda.")
setph(s,1,"Kullan → Üret → Standartlaştır → Otomatikleştir   ·   © Fikri Eren 2026")

pr.save(OUT)
print("KAYDEDILDI:",OUT,"| slayt:",len(pr.slides._sldIdLst))
